"""Extrai texto de arquivos .pptx e estrutura em fichas/experimentos.

Uso:
    python -m metaexp extrair-pptx <caminho-do-zip-ou-pasta>

Descompacta o ZIP (se fornecido), varre todos os .pptx encontrados, extrai o
texto de cada slide e usa o LLM para estruturar o conteúdo nos schemas Ficha
e Experimento. Os resultados são salvos em data/corpus/real/.
"""

from __future__ import annotations

import logging
import re
import shutil
import tempfile
import time
import zipfile
from pathlib import Path

from pptx import Presentation

from .config import settings
from .context.builder import system_ingestao
from .corpus.store import save
from .llm.client import AnthropicLLM, LLM
from .schemas import Experimento, Ficha

log = logging.getLogger("metaexp.extract_pptx")

SUPORTADOS = {".pptx"}
MAX_RETRIES = 3
RETRY_DELAY = 5  # segundos entre retentativas


def _extrair_texto_pptx(path: Path) -> str:
    """Extrai todo o texto de um arquivo .pptx, slide a slide."""
    prs = Presentation(str(path))
    partes: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        textos_slide: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    texto = paragraph.text.strip()
                    if texto:
                        textos_slide.append(texto)
            if shape.has_table:
                tabela = shape.table
                for row in tabela.rows:
                    celulas = [cell.text.strip() for cell in row.cells]
                    textos_slide.append(" | ".join(celulas))
        if textos_slide:
            partes.append(f"--- Slide {i} ---")
            partes.extend(textos_slide)
    return "\n".join(partes)


def _slug(nome: str) -> str:
    """Gera um identificador a partir do nome do arquivo."""
    base = Path(nome).stem
    base = re.sub(r"[^A-Za-z0-9]+", "-", base).strip("-").upper()[:40]
    return base if base else "EXPERIMENTO"


def _chamar_com_retry(fn, *args, **kwargs):
    """Chama fn com retry em caso de erro de conexão."""
    ultimo_erro = None
    for tentativa in range(1, MAX_RETRIES + 1):
        try:
            return fn(*args, **kwargs)
        except Exception as e:
            erro_msg = str(e)
            # Só retenta em erros de conexão/rede
            if any(palavra in erro_msg.lower() for palavra in
                   ["connect", "read", "timeout", "reset", "refused", "unreachable",
                    "10065", "10054", "10060", "10061"]):
                ultimo_erro = e
                if tentativa < MAX_RETRIES:
                    log.warning("  erro de conexão (tentativa %d/%d): %s. Aguardando %ds...",
                                tentativa, MAX_RETRIES, e, RETRY_DELAY)
                    time.sleep(RETRY_DELAY)
                else:
                    log.error("  falha após %d tentativas: %s", MAX_RETRIES, e)
                    raise
            else:
                # Erro não é de conexão, propaga imediatamente
                raise
    raise ultimo_erro  # type: ignore[misc]


def _extrair_ficha_do_texto(llm: LLM, texto: str, uid: str) -> Ficha:
    """Usa o LLM para estruturar o texto extraído em uma Ficha."""
    system = system_ingestao()
    prompt = (
        f"O texto abaixo foi extraído de uma apresentação PowerPoint sobre um experimento de IA/Dados "
        f"da beOn Labs. Extraia dele uma Ficha de Experimentação completa no schema Ficha.\n\n"
        f"ID do experimento: REAL-{uid}\n"
        f"Origem: real\n\n"
        f"Texto extraído:\n\n{texto}\n\n"
        f"Preencha todos os campos da Ficha que conseguir inferir do texto. "
        f"Para campos obrigatórios sem informação no texto, deixe como null. "
        f"Use 'outro' para domínio/tecnologia quando não for possível determinar."
    )
    return _chamar_com_retry(
        llm.structured,
        system=system,
        messages=[{"role": "user", "content": [{"type": "text", "text": prompt}]}],
        schema=Ficha,
        effort=settings.effort_synth,
        max_tokens=16000,
    )


def _extrair_experimento_do_texto(llm: LLM, texto: str, uid: str, ficha: Ficha) -> Experimento:
    """Usa o LLM para estruturar o texto completo em um Experimento."""
    system = system_ingestao()
    prompt = (
        f"O texto abaixo foi extraído de uma apresentação PowerPoint sobre um experimento de IA/Dados "
        f"da beOn Labs. Converta-o em um registro completo de Experimento.\n\n"
        f"ID: REAL-{uid}\n"
        f"Origem: real\n\n"
        f"A Ficha já extraída é:\n{ficha.model_dump_json(indent=2)}\n\n"
        f"Texto extraído:\n\n{texto}\n\n"
        f"Preencha todos os campos do Experimento que conseguir inferir. "
        f"Inclua descricao_amostra, e se houver informações sobre resultados ou "
        f"parecer, preencha também resultados e parecer. "
        f"Use 'outro' para domínio quando não for possível determinar."
    )
    return _chamar_com_retry(
        llm.structured,
        system=system,
        messages=[{"role": "user", "content": [{"type": "text", "text": prompt}]}],
        schema=Experimento,
        effort=settings.effort_synth,
        max_tokens=32000,
    )


def _eh_template(nome: str) -> bool:
    """Verifica se o arquivo é um template (começa com [Template])."""
    return nome.strip().startswith("[Template]") or nome.strip().startswith("[template]")


def processar_pptx(llm: LLM, pptx_path: Path, destino: Path, idx: int, total: int) -> str | None:
    """Processa um único .pptx: extrai texto, estrutura e salva."""
    log.info("[%d/%d] %s", idx, total, pptx_path.name)

    # Pula templates
    if _eh_template(pptx_path.stem):
        log.info("  pulando template")
        return None

    texto = _extrair_texto_pptx(pptx_path)
    if not texto.strip():
        log.warning("  texto vazio, ignorando")
        return None

    uid = _slug(pptx_path.stem)

    # Primeiro extrai a ficha
    ficha = _extrair_ficha_do_texto(llm, texto, uid)
    ficha.id = f"REAL-{uid}"

    # Depois monta o experimento completo
    exp = _extrair_experimento_do_texto(llm, texto, uid, ficha)
    exp.id = f"REAL-{uid}"
    exp.origem = "real"
    exp.ficha.id = exp.id

    save(exp, destino)
    log.info("  OK -> %s.json", exp.id)
    return exp.id


def extrair_pptx(
    origem: Path,
    destino: Path | None = None,
    llm: LLM | None = None,
) -> list[str]:
    """Extrai texto de .pptx e estrutura em experimentos.

    Args:
        origem: Caminho para o arquivo .zip ou para uma pasta com .pptx.
        destino: Pasta onde salvar os JSONs (default: data/corpus/real/).
        llm: Instância do LLM (cria uma nova se não fornecida).

    Returns:
        Lista de IDs dos experimentos gerados.
    """
    destino = destino or settings.corpus_dirs[0]
    llm = llm or AnthropicLLM()

    # Se for um ZIP, descompacta para pasta temporária
    if origem.suffix.lower() == ".zip":
        tmp = Path(tempfile.mkdtemp(prefix="metaexp-pptx-"))
        try:
            log.info("descompactando %s em %s", origem, tmp)
            with zipfile.ZipFile(str(origem), "r") as zf:
                zf.extractall(str(tmp))
            return _processar_pasta(tmp, destino, llm)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    elif origem.is_dir():
        return _processar_pasta(origem, destino, llm)
    else:
        log.error("origem deve ser um .zip ou uma pasta: %s", origem)
        return []


def _processar_pasta(pasta: Path, destino: Path, llm: LLM) -> list[str]:
    """Varre uma pasta em busca de .pptx e processa cada um."""
    pptxs = sorted(pasta.rglob("*.pptx"))
    if not pptxs:
        log.warning("nenhum arquivo .pptx encontrado em %s", pasta)
        return []

    # Filtra templates
    pptxs = [p for p in pptxs if not _eh_template(p.stem)]
    total = len(pptxs)
    log.info("encontrados %d arquivo(s) .pptx (após filtrar templates)", total)

    gerados: list[str] = []
    for idx, pptx_path in enumerate(pptxs, 1):
        try:
            exp_id = processar_pptx(llm, pptx_path, destino, idx, total)
            if exp_id:
                gerados.append(exp_id)
        except Exception:
            log.exception("  erro ao processar (pulado)")
    log.info("concluído: %d/%d processados com sucesso", len(gerados), total)
    return gerados