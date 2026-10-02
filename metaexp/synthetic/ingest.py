"""Ingestão de documentos reais (fichas, relatórios, pareceres) para o corpus.

Coloque os arquivos em data/real/ (essa pasta não vai para o Git). Formatos:
  .pdf          enviado ao modelo como documento (texto e layout)
  .docx         texto e tabelas extraídos com python-docx
  .md, .txt     lidos como texto
  .json         já no formato Experimento: validado e copiado
Vários arquivos do mesmo experimento podem ir numa subpasta (ex.: ficha.docx e
relatorio.pdf em data/real/EXP-VOC-01/): eles viram um único registro.
"""

from __future__ import annotations

import base64
import logging
import re
from pathlib import Path

from ..config import Settings, settings as default_settings
from ..context.builder import system_ingestao
from ..corpus.store import save
from ..llm.client import LLM
from ..schemas import Experimento

log = logging.getLogger("metaexp.ingest")
SUPORTADOS = {".pdf", ".docx", ".md", ".txt", ".json"}


def _docx_text(path: Path) -> str:
    import docx  # python-docx

    d = docx.Document(str(path))
    partes = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            partes.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(partes)


def _blocos(path: Path) -> list[dict]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        data = base64.standard_b64encode(path.read_bytes()).decode("ascii")
        return [{"type": "document", "title": path.name,
                 "source": {"type": "base64", "media_type": "application/pdf", "data": data}}]
    if ext == ".docx":
        texto = _docx_text(path)
    else:
        texto = path.read_text(encoding="utf-8", errors="replace")
    return [{"type": "text", "text": f"<documento nome=\"{path.name}\">\n{texto}\n</documento>"}]


def _slug(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-").upper()[:40]


def unidades(origem: Path) -> list[tuple[str, list[Path]]]:
    """Agrupa arquivos por experimento: cada subpasta ou cada arquivo solto é uma unidade."""
    out: list[tuple[str, list[Path]]] = []
    for item in sorted(origem.iterdir()):
        if item.name.startswith(".") or item.name.lower() == "readme.md":
            continue
        if item.is_dir():
            files = [f for f in sorted(item.rglob("*")) if f.suffix.lower() in SUPORTADOS]
            if files:
                out.append((_slug(item.name), files))
        elif item.suffix.lower() in SUPORTADOS:
            out.append((_slug(item.stem), [item]))
    return out


def ingerir(llm: LLM, origem: Path | None = None, destino: Path | None = None,
            cfg: Settings | None = None) -> list[str]:
    cfg = cfg or default_settings
    origem = origem or cfg.real_docs_dir
    destino = destino or (cfg.corpus_dirs[0] if cfg.corpus_dirs else origem.parent / "corpus/real")
    system = system_ingestao()
    gerados = []
    for uid, files in unidades(origem):
        if len(files) == 1 and files[0].suffix.lower() == ".json":
            exp = Experimento.model_validate_json(files[0].read_text(encoding="utf-8"))
        else:
            content: list[dict] = []
            for f in files:
                content += _blocos(f)
            content.append({"type": "text", "text":
                            f"Converta os documentos acima em um único registro de experimento. Use o id REAL-{uid} e origem 'real'."})
            exp = llm.structured(system=system, messages=[{"role": "user", "content": content}],
                                 schema=Experimento, effort=cfg.effort_synth, max_tokens=32000)
        exp.id, exp.origem = f"REAL-{uid}" if not exp.id.startswith("REAL-") else exp.id, "real"
        exp.ficha.id = exp.id
        save(exp, destino)
        gerados.append(exp.id)
        log.info("ingerido %s (%d arquivo(s))", exp.id, len(files))
    return gerados
