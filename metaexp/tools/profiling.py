"""Perfil determinístico de arquivos de amostra enviados no chat.

O modelo não lê o arquivo inteiro: recebe este perfil (estrutura, volume,
qualidade e algumas linhas de exemplo). Assim a análise é barata, reprodutível
e não depende do tamanho do arquivo.
"""

from __future__ import annotations

import csv
import io
import json
from dataclasses import asdict, dataclass, field
from pathlib import Path


@dataclass
class PerfilArquivo:
    nome: str
    formato: str
    registros: int
    colunas: list[str] = field(default_factory=list)
    completude: float = 1.0           # fração de células preenchidas
    duplicadas: int = 0               # linhas idênticas repetidas
    vazias_por_coluna: dict[str, int] = field(default_factory=dict)
    exemplos: list[dict] = field(default_factory=list)
    observacoes: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return asdict(self)


def _profile_rows(nome: str, formato: str, header: list[str], rows: list[list[str]]) -> PerfilArquivo:
    header = [h.strip() or f"coluna_{i + 1}" for i, h in enumerate(header)]
    total_cells = len(rows) * max(len(header), 1)
    vazias = {h: 0 for h in header}
    filled = 0
    for r in rows:
        for i, h in enumerate(header):
            v = r[i] if i < len(r) else ""
            if str(v).strip():
                filled += 1
            else:
                vazias[h] += 1
    seen, dup = set(), 0
    for r in rows:
        key = tuple(str(x).strip().lower() for x in r)
        if key in seen:
            dup += 1
        seen.add(key)
    exemplos = [{h: (str(r[i])[:200] if i < len(r) else "") for i, h in enumerate(header)} for r in rows[:5]]
    p = PerfilArquivo(nome=nome, formato=formato, registros=len(rows), colunas=header,
                      completude=round(filled / total_cells, 4) if total_cells else 0.0,
                      duplicadas=dup, vazias_por_coluna={k: v for k, v in vazias.items() if v}, exemplos=exemplos)
    if not rows:
        p.observacoes.append("arquivo sem registros")
    return p


def _from_csv(nome: str, data: bytes) -> PerfilArquivo:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    if not rows:
        return _profile_rows(nome, "csv", [], [])
    return _profile_rows(nome, "csv", rows[0], rows[1:])


def _from_xlsx(nome: str, data: bytes) -> PerfilArquivo:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(c.strip() for c in r)]
    p = _profile_rows(nome, "xlsx", rows[0] if rows else [], rows[1:])
    if len(wb.worksheets) > 1:
        p.observacoes.append(f"planilha com {len(wb.worksheets)} abas; analisada apenas a primeira ({ws.title})")
    return p


def _from_json(nome: str, data: bytes) -> PerfilArquivo:
    obj = json.loads(data.decode("utf-8-sig"))
    if isinstance(obj, dict):
        obj = next((v for v in obj.values() if isinstance(v, list)), [obj])
    if not isinstance(obj, list):
        return _profile_rows(nome, "json", ["valor"], [[str(obj)]])
    header: list[str] = []
    for item in obj:
        if isinstance(item, dict):
            for k in item:
                if k not in header:
                    header.append(k)
    rows = [[str(item.get(h, "")) if isinstance(item, dict) else str(item) for h in header] for item in obj]
    return _profile_rows(nome, "json", header or ["valor"], rows)


def _from_text(nome: str, data: bytes) -> PerfilArquivo:
    text = data.decode("utf-8", errors="replace")
    blocos = [b.strip() for b in text.split("\n\n") if b.strip()]
    p = _profile_rows(nome, "texto", ["trecho"], [[b] for b in blocos])
    p.observacoes.append("texto livre: cada parágrafo foi contado como um registro")
    return p


def perfilar(nome: str, data: bytes) -> PerfilArquivo:
    ext = Path(nome).suffix.lower()
    if ext in (".csv", ".tsv"):
        return _from_csv(nome, data)
    if ext in (".xlsx", ".xlsm"):
        return _from_xlsx(nome, data)
    if ext == ".json":
        return _from_json(nome, data)
    if ext in (".txt", ".md"):
        return _from_text(nome, data)
    p = PerfilArquivo(nome=nome, formato=ext.lstrip(".") or "desconhecido", registros=0)
    p.observacoes.append("formato não suportado para perfil automático; peça CSV, XLSX, JSON ou TXT")
    return p
